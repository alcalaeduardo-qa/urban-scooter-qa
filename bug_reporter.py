"""
bug_reporter.py — Genera reporte de bugs en formato Jira (Excel + CSV).

Escanea los reportes más recientes en results/, extrae filas "No Aprobado"
y produce:
  - bugs_jira_<ts>.xlsx / .csv  → listos para importar a Jira
  - bug_attachments_<ts>/       → carpeta con screenshots de cada bug
  - bug_manifest_<ts>.json      → mapeo Resumen → ruta screenshot (para jira_uploader.py)
"""
from __future__ import annotations

import csv
import json
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl import Workbook

ROOT            = Path(__file__).resolve().parent
RESULTS_DIR     = ROOT / "results"
SCREENSHOTS_DIR = ROOT / "screenshots"

JIRA_HEADERS = [
    "Resumen",
    "Clave de incidencia",
    "Tipo de Incidencia",
    "Estado",
    "Clave del proyecto",
    "Nombre del proyecto",
    "Tipo de proyecto",
    "Prioridad",
    "Descripción",
    "Entorno",
]

# ── Configuración de cada tarea ────────────────────────────────────────────────
# prefix          : coincide con el inicio del nombre del archivo en results/
# task_label      : etiqueta legible para el campo Resumen
# sheet_index     : hoja del Excel (0 = primera)
# env             : valor para columna Entorno en Jira
# status_header   : texto del encabezado de la columna Estado
# id_col          : encabezado de la columna ID/caso
# desc_col        : encabezado de la columna descripción principal
# extra_cols      : columnas adicionales útiles para la descripción Jira
# screenshot_col  : índice 0-based de la columna donde write_status guarda la ruta
#                   del screenshot (None = no aplica para esta tarea)
# screenshot_dir  : subdirectorio dentro de screenshots/ para esta tarea
TASK_CONFIGS: list[dict[str, Any]] = [
    {
        "prefix":         "Tarea_2_lista",
        "task_label":     "T2-Checklist-Opera",
        "sheet_index":    0,
        "env":            "Web — Opera",
        "status_header":  "Opera",
        "id_col":         "ID de Comprobación",
        "desc_col":       "Descripción (Acción/Verbo)",
        "extra_cols":     [],
        "screenshot_col": None,
        "screenshot_dir": "tarea2_checklist",
    },
    {
        "prefix":         "Tarea_2_lista",
        "task_label":     "T2-Checklist-Chrome",
        "sheet_index":    0,
        "env":            "Web — Chrome",
        "status_header":  "Chrome",
        "id_col":         "ID de Comprobación",
        "desc_col":       "Descripción (Acción/Verbo)",
        "extra_cols":     [],
        "screenshot_col": None,
        "screenshot_dir": "tarea2_checklist",
    },
    {
        "prefix":         "Tarea_2_validaci",
        "task_label":     "T2-Validación-Opera",
        "sheet_index":    0,
        "env":            "Web — Opera",
        "status_header":  "Opera",
        "id_col":         "ID",
        "desc_col":       "Campo",
        "extra_cols":     ["Nombre de la clase", "Límites", "Explicación"],
        "screenshot_col": None,
        "screenshot_dir": "tarea_2_validacion",
    },
    {
        "prefix":         "Tarea_2_validaci",
        "task_label":     "T2-Validación-Chrome",
        "sheet_index":    0,
        "env":            "Web — Chrome",
        "status_header":  "Chrome",
        "id_col":         "ID",
        "desc_col":       "Campo",
        "extra_cols":     ["Nombre de la clase", "Límites", "Explicación"],
        "screenshot_col": None,
        "screenshot_dir": "tarea_2_validacion",
    },
    {
        "prefix":         "Tarea_3_casos",
        "task_label":     "T3-Casos",
        "sheet_index":    0,
        "env":            "Android 13 — emulator-5554 — App v1.0",
        "status_header":  "Estado (Aprobado/No Aprobado)",
        "id_col":         "ID del caso de prueba",
        "desc_col":       "Nombre del caso de prueba",
        "extra_cols":     ["Descripción de los pasos", "Resultado esperado", "Comentarios"],
        "screenshot_col": 7,   # col H — write_status hardcodea column=8 (1-indexed)
        "screenshot_dir": "tarea3_mobile",
    },
    {
        "prefix":         "Tarea_4_lista",
        "task_label":     "T4-API",
        "sheet_index":    0,
        "env":            "API REST — servidor TripleTen",
        "status_header":  "Estado (Aprobado/No Aprobado)",
        "id_col":         "ID",
        "desc_col":       "Descripción de la prueba (Paso a comprobar)",
        "extra_cols":     ["Endpoint (Bloque de Pruebas)", "Código de Respuesta Esperado"],
        "screenshot_col": 7,   # col H — write_status hardcodea column=8 (1-indexed)
        "screenshot_dir": "tarea4_api",
    },
    {
        "prefix":         "Tarea_5_happy",
        "task_label":     "T5-HappyPath-Opera",
        "sheet_index":    0,
        "env":            "Opera — Android 13 — emulator-5554",
        "status_header":  "Estado",
        "id_col":         "ID del caso",
        "desc_col":       "Nombre del caso",
        "extra_cols":     ["Pasos", "Resultado esperado", "Comentarios"],
        "screenshot_col": 7,   # col H — write_status hardcodea column=8 (1-indexed)
        "screenshot_dir": "tarea5",
    },
    {
        "prefix":         "Tarea_5_happy",
        "task_label":     "T5-HappyPath-Chrome",
        "sheet_index":    1,
        "env":            "Chrome (Opera fallback) — Android 13 — emulator-5554",
        "status_header":  "Estado",
        "id_col":         "ID del caso",
        "desc_col":       "Nombre del caso",
        "extra_cols":     ["Pasos", "Resultado esperado", "Comentarios"],
        "screenshot_col": 7,   # col H — write_status hardcodea column=8 (1-indexed)
        "screenshot_dir": "tarea5",
    },
]


# ══════════════════════════════════════════════════════════════════════════════
# Helpers
# ══════════════════════════════════════════════════════════════════════════════

def _latest_result(prefix: str) -> Path | None:
    """Devuelve el archivo más reciente en results/ que empiece con prefix."""
    candidates = sorted(
        RESULTS_DIR.glob(f"{prefix}*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    return candidates[0] if candidates else None


def _col_index(headers: list[str], header_name: str) -> int | None:
    """Devuelve el índice (0-based) de la columna cuyo encabezado coincida."""
    for i, h in enumerate(headers):
        if h and header_name.strip().lower() in str(h).strip().lower():
            return i
    return None


def _build_description(row: tuple, col_map: dict[str, int], extra_cols: list[str]) -> str:
    """Construye el campo Descripción en formato Jira."""
    parts: list[str] = []

    # Descripción principal desde los extras
    for col_name in extra_cols:
        idx = col_map.get(col_name)
        if idx is not None and idx < len(row) and row[idx]:
            val = str(row[idx]).strip()
            if val and val.lower() != "none":
                parts.append(f"• *{col_name}:* {val}")

    desc = "\n".join(parts) if parts else "Ver captura de pantalla adjunta."

    return (
        f"* 📃*Descripción:* {desc}\n"
        f"* 📌*Pasos para reproducir:* Ver columna 'Pasos' del reporte Excel.\n"
        f"* 🔴*Resultado actual:* No Aprobado — el test falló durante la ejecución automatizada.\n"
        f"* 🟢*Resultado esperado:* El test debe pasar sin errores.\n"
    )


def _extract_bugs(config: dict[str, Any]) -> list[dict[str, str]]:
    """Lee el último resultado de una tarea y extrae los bugs (No Aprobado)."""
    result_file = _latest_result(config["prefix"])
    if not result_file:
        return []

    wb = openpyxl.load_workbook(result_file, data_only=True)
    try:
        ws = wb.worksheets[config["sheet_index"]]
    except IndexError:
        return []

    # Leer encabezados
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_map = {h: i for i, h in enumerate(headers) if h}

    status_idx = _col_index(headers, config["status_header"])
    id_idx     = _col_index(headers, config["id_col"])
    desc_idx   = _col_index(headers, config["desc_col"])

    if status_idx is None:
        return []

    bugs: list[dict[str, str]] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        status_val = str(row[status_idx] or "").strip() if status_idx < len(row) else ""
        if "no aprobado" not in status_val.lower():
            continue

        case_id  = str(row[id_idx] or "").strip()  if id_idx  is not None and id_idx  < len(row) else ""
        desc_val = str(row[desc_idx] or "").strip() if desc_idx is not None and desc_idx < len(row) else ""

        if not case_id and not desc_val:
            continue

        resumen = f"[{config['task_label']}] [{case_id}] {desc_val}"
        jira_desc = _build_description(row, col_map, config.get("extra_cols", []))

        # Leer ruta de screenshot si la tarea la guarda en el Excel
        screenshot_path = ""
        sc_col = config.get("screenshot_col")
        if sc_col is not None and sc_col < len(row) and row[sc_col]:
            candidate = str(row[sc_col]).strip()
            if candidate and Path(candidate).exists():
                screenshot_path = candidate

        bugs.append({
            "Resumen":            resumen[:255],
            "Clave de incidencia": "",
            "Tipo de Incidencia":  "Error",
            "Estado":              "Tareas por hacer",
            "Clave del proyecto":  "PFTT",
            "Nombre del proyecto": "Urban_Scooter",
            "Tipo de proyecto":    "software",
            "Prioridad":           "Medium",
            "Descripción":         jira_desc,
            "Entorno":             config["env"],
            "_screenshot":         screenshot_path,   # campo interno, no va al CSV Jira
            "_screenshot_dir":     config.get("screenshot_dir", ""),
        })

    return bugs


# ══════════════════════════════════════════════════════════════════════════════
# Generadores de archivos
# ══════════════════════════════════════════════════════════════════════════════

def _write_excel(bugs: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bugs Jira"

    # Encabezados
    header_fill = PatternFill(start_color="7C6AF7", end_color="7C6AF7", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, header in enumerate(JIRA_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill  = header_fill
        cell.font  = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Filas de datos
    red_fill   = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    for row_idx, bug in enumerate(bugs, start=2):
        for col_idx, header in enumerate(JIRA_HEADERS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=bug.get(header, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if header == "Tipo de Incidencia":
                cell.fill = red_fill

    # Anchos de columna
    col_widths = [60, 20, 18, 18, 18, 18, 14, 12, 70, 40]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    ws.row_dimensions[1].height = 30
    wb.save(path)


def _write_csv(bugs: list[dict], path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=JIRA_HEADERS)
        writer.writeheader()
        writer.writerows(bugs)


# ══════════════════════════════════════════════════════════════════════════════
# Attachments
# ══════════════════════════════════════════════════════════════════════════════

def _safe_filename(text: str, max_len: int = 80) -> str:
    """Convierte un texto en nombre de archivo seguro."""
    safe = re.sub(r"[^\w\-]", "_", text)
    return safe[:max_len]


def _collect_attachments(bugs: list[dict], attachments_dir: Path) -> list[dict]:
    """
    Para cada bug:
      1. Si tiene screenshot_path → copia el archivo a attachments_dir.
      2. Si no tiene pero screenshot_dir existe → busca el screenshot más reciente
         en screenshots/<screenshot_dir>/ y lo copia.
    Devuelve la lista de bugs con el campo '_attachment' actualizado.
    """
    attachments_dir.mkdir(parents=True, exist_ok=True)

    for bug in bugs:
        src: Path | None = None

        # Prioridad 1: ruta exacta guardada en el Excel
        if bug.get("_screenshot") and Path(bug["_screenshot"]).exists():
            src = Path(bug["_screenshot"])
        else:
            # Prioridad 2: screenshot más reciente del directorio de la tarea
            sc_dir = SCREENSHOTS_DIR / bug.get("_screenshot_dir", "")
            if sc_dir.is_dir():
                candidates = sorted(
                    list(sc_dir.glob("*.png")) + list(sc_dir.glob("*.txt")),
                    key=lambda p: p.stat().st_mtime,
                    reverse=True,
                )
                # Intentar filtrar por ID/case si está en el Resumen
                resumen = bug.get("Resumen", "")
                matched = [
                    p for p in candidates
                    if any(part in p.stem for part in resumen.replace("]", "[").split("[") if part.strip())
                ]
                src = matched[0] if matched else (candidates[0] if candidates else None)

        bug["_attachment"] = ""
        if src:
            dest_name = f"{_safe_filename(bug['Resumen'])}{src.suffix}"
            dest = attachments_dir / dest_name
            shutil.copy2(src, dest)
            bug["_attachment"] = str(dest)

    return bugs


def _write_manifest(bugs: list[dict], path: Path) -> None:
    """
    Escribe bug_manifest_<ts>.json:
    Lista de objetos con Resumen y ruta del attachment (para jira_uploader.py).
    """
    manifest = [
        {
            "resumen":    bug["Resumen"],
            "attachment": bug.get("_attachment", ""),
        }
        for bug in bugs
    ]
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


# ══════════════════════════════════════════════════════════════════════════════
# Punto de entrada público
# ══════════════════════════════════════════════════════════════════════════════

def generate_bug_report() -> tuple[Path, Path, Path, Path, int]:
    """
    Extrae bugs, copia screenshots y genera:
      - bugs_jira_<ts>.xlsx / .csv  → para importar a Jira
      - bug_attachments_<ts>/       → carpeta con screenshots
      - bug_manifest_<ts>.json      → mapeo Resumen → attachment (para jira_uploader.py)

    Devuelve (excel_path, csv_path, attachments_dir, manifest_path, total_bugs).
    """
    all_bugs: list[dict] = []
    seen: set[str] = set()

    for config in TASK_CONFIGS:
        for bug in _extract_bugs(config):
            key = bug["Resumen"]
            if key not in seen:
                seen.add(key)
                all_bugs.append(bug)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_path       = RESULTS_DIR / f"bugs_jira_{ts}.xlsx"
    csv_path        = RESULTS_DIR / f"bugs_jira_{ts}.csv"
    attachments_dir = RESULTS_DIR / f"bug_attachments_{ts}"
    manifest_path   = RESULTS_DIR / f"bug_manifest_{ts}.json"

    # Recopilar screenshots → attachments_dir
    all_bugs = _collect_attachments(all_bugs, attachments_dir)

    # Escribir manifest JSON
    _write_manifest(all_bugs, manifest_path)

    # Limpiar campos internos antes de escribir CSV/Excel
    jira_bugs = [{k: v for k, v in bug.items() if not k.startswith("_")} for bug in all_bugs]

    _write_excel(jira_bugs, xlsx_path)
    _write_csv(jira_bugs, csv_path)

    return xlsx_path, csv_path, attachments_dir, manifest_path, len(all_bugs)
